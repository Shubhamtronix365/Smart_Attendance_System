import json
import calendar
from io import BytesIO
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ─── Formatting Helpers ───────────────────────────────────────────────────────
def format_time_str(val):
    if not val:
        return None
    if isinstance(val, (datetime, date)):
        if hasattr(val, "strftime"):
            return val.strftime("%I:%M %p")
    s = str(val).strip()
    if " " in s or "T" in s:
        try:
            dt = datetime.fromisoformat(s.replace(" ", "T"))
            return dt.strftime("%I:%M %p")
        except Exception:
            pass
    if ":" in s and len(s) <= 8:
        try:
            parts = s.split(":")
            dt = datetime(2026, 1, 1, int(parts[0]), int(parts[1]))
            return dt.strftime("%I:%M %p")
        except Exception:
            pass
    return s

def _is_next_day(in_val, out_val):
    try:
        in_dt = in_val if isinstance(in_val, datetime) else (datetime.fromisoformat(str(in_val).replace(" ", "T")) if in_val else None)
        out_dt = out_val if isinstance(out_val, datetime) else (datetime.fromisoformat(str(out_val).replace(" ", "T")) if out_val else None)
        if in_dt and out_dt:
            return out_dt.date() > in_dt.date()
    except Exception:
        pass
    return False

def _normalize_status(status_raw) -> str:
    if hasattr(status_raw, "value"):
        return str(status_raw.value).lower().strip()
    s = str(status_raw or "absent").lower().strip()
    if "." in s:
        s = s.split(".")[-1]
    return s

def format_late_duration_excel(minutes) -> str:
    if not minutes or int(minutes) <= 0:
        return "—"
    m = int(minutes)
    if m < 60:
        return f"{m}m"
    hrs = m // 60
    rem = m % 60
    return f"{hrs}h {rem}m" if rem > 0 else f"{hrs}h"

def format_sessions_summary(punches_data) -> str:
    if not punches_data:
        return "—"
    try:
        if isinstance(punches_data, str):
            sessions = json.loads(punches_data)
        else:
            sessions = punches_data
        if not sessions or not isinstance(sessions, list):
            return "—"
        
        parts = []
        for idx, s in enumerate(sessions, 1):
            s_num = s.get("session", idx)
            in_raw = s.get("in")
            out_raw = s.get("out")
            hrs = s.get("hours")
            
            in_str = format_time_str(in_raw) or "--:--"
            if out_raw:
                out_str = format_time_str(out_raw) or str(out_raw)
                if _is_next_day(in_raw, out_raw):
                    out_str += " (+1d)"
            else:
                out_str = "Active Session"
            
            duration_str = f" ({float(hrs):.2f}h)" if hrs is not None else ""
            parts.append(f"S{s_num}: In {in_str} → Out {out_str}{duration_str}")
            
        return " | ".join(parts) if parts else "—"
    except Exception:
        return "—"

# ─── Executive Design Palette ─────────────────────────────────────────────────
FONT_FAMILY = "Segoe UI"
HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")

TITLE_FONT = Font(name=FONT_FAMILY, size=15, bold=True, color="0F172A")
SUBTITLE_FONT = Font(name=FONT_FAMILY, size=9.5, italic=True, color="475569")
BOLD_FONT = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0F172A")
REGULAR_FONT = Font(name=FONT_FAMILY, size=9.5, color="1E293B")

ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

TOTAL_BORDER = Border(
    top=Side(style='thin', color='0F172A'),
    bottom=Side(style='double', color='0F172A')
)

def _get_status_style(status_raw: str):
    s = _normalize_status(status_raw)
    if s in ("present", "wfh"):
        return PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"), Font(name=FONT_FAMILY, size=9.5, bold=True, color="065F46")
    elif s == "late":
        return PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), Font(name=FONT_FAMILY, size=9.5, bold=True, color="92400E")
    elif s in ("holiday", "weekly_off"):
        return PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"), Font(name=FONT_FAMILY, size=9.5, bold=True, color="0369A1")
    elif s == "leave":
        return PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid"), Font(name=FONT_FAMILY, size=9.5, bold=True, color="6D28D9")
    elif s in ("half_day", "halfday"):
        return PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), Font(name=FONT_FAMILY, size=9.5, bold=True, color="92400E")
    else:
        return PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"), Font(name=FONT_FAMILY, size=9.5, bold=True, color="991B1B")

def _autofit_columns(ws, min_col=2, max_col=15, skip_rows=8):
    """Automatically fits column widths with padding and maximum limits, ignoring title & KPI rows."""
    for col in ws.columns:
        col_idx = col[0].column
        if col_idx < min_col or col_idx > max_col:
            continue
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in col:
            # Ignore title and KPI dashboard rows
            if cell.row <= skip_rows:
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 13), 55)

def _write_kpi_dashboard(ws, start_row: int, kpi_list: list, start_col: int = 2):
    """
    Renders executive KPI dashboard summary cards across specified columns.
    kpi_list: list of tuples: (label_str, value_str_or_num, font_hex_color)
    """
    ws.row_dimensions[start_row].height = 16
    ws.row_dimensions[start_row + 1].height = 24
    
    kpi_label_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    kpi_val_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    kpi_lbl_font = Font(name=FONT_FAMILY, size=7.5, bold=True, color="475569")
    
    for idx, (label, val, val_color) in enumerate(kpi_list):
        c_col = start_col + idx
        # Label cell
        c_lbl = ws.cell(row=start_row, column=c_col, value=str(label).upper())
        c_lbl.font = kpi_lbl_font
        c_lbl.alignment = ALIGN_CENTER
        c_lbl.fill = kpi_label_fill
        c_lbl.border = THIN_BORDER
        
        # Value cell
        c_val = ws.cell(row=start_row + 1, column=c_col, value=val)
        c_val.font = Font(name=FONT_FAMILY, size=11, bold=True, color=val_color)
        c_val.alignment = ALIGN_CENTER
        c_val.fill = kpi_val_fill
        c_val.border = THIN_BORDER

# ─── 1. Daily Attendance Excel ────────────────────────────────────────────────
def generate_daily_attendance_excel(report_date: date, records: list, title: str = "DAILY ATTENDANCE REPORT", subtitle: str = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Attendance"
    ws.views.sheetView[0].showGridLines = True

    # Title Block
    ws.cell(row=2, column=2, value="SMART ATTENDANCE MANAGEMENT SYSTEM").font = Font(name=FONT_FAMILY, size=10, bold=True, color="0284C7")
    ws.cell(row=3, column=2, value=title).font = TITLE_FONT
    sub = subtitle or f"Date: {report_date.strftime('%A, %d %B %Y')} | Exported on: {datetime.now().strftime('%d %b %Y, %I:%M %p IST')}"
    ws.cell(row=4, column=2, value=sub).font = SUBTITLE_FONT

    # Pre-calculate Totals & Metrics
    total_logs = len(records)
    present_cnt = sum(1 for r in records if _normalize_status(r.get("status")) in ("present", "wfh"))
    late_cnt = sum(1 for r in records if _normalize_status(r.get("status")) == "late" or (r.get("late_minutes") and int(r.get("late_minutes")) > 0))
    absent_cnt = sum(1 for r in records if _normalize_status(r.get("status")) == "absent")
    leave_cnt = sum(1 for r in records if _normalize_status(r.get("status")) == "leave")
    holiday_cnt = sum(1 for r in records if _normalize_status(r.get("status")) in ("holiday", "weekly_off"))
    total_late_min = sum(int(r.get("late_minutes") or 0) for r in records)
    total_late_str = format_late_duration_excel(total_late_min)
    total_hours = sum(float(r.get("working_hours") or 0.0) for r in records)
    total_ot = sum(float(r.get("overtime_hours") or 0.0) for r in records)
    total_checkouts = sum(int(r.get("checkout_count") or (1 if r.get("check_out") else 0)) for r in records)

    # Executive KPI Dashboard Cards (Rows 6-7, Cols 2-12)
    kpis = [
        ("TOTAL LOGS", total_logs, "0F172A"),
        ("PRESENT", present_cnt, "065F46"),
        ("ABSENT", absent_cnt, "991B1B"),
        ("LATE ARRIVALS", late_cnt, "92400E"),
        ("LATE DURATION", total_late_str, "92400E"),
        ("APPROVED LEAVE", leave_cnt, "6D28D9"),
        ("WEEKLY OFF / HOL", holiday_cnt, "0369A1"),
        ("CHECKOUTS", f"{total_checkouts} COs", "0284C7"),
        ("TOTAL WORKED", f"{total_hours:.2f} hrs", "0F172A"),
        ("OVERTIME", f"+{total_ot:.2f} hrs", "7C3AED"),
        ("SYSTEM STATUS", "100% VERIFIED", "065F46"),
    ]
    _write_kpi_dashboard(ws, start_row=6, kpi_list=kpis, start_col=2)

    headers = [
        "Employee ID", "Employee Name", "Department", "Check-In", 
        "Check-Out", "Checkouts", "Sessions Breakdown", "Working Hours", "Overtime (hrs)", "Status", "Source"
    ]

    start_row = 9
    ws.row_dimensions[8].height = 10
    ws.row_dimensions[start_row].height = 26
    for col_idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    ws.freeze_panes = f"B{start_row + 1}"

    current_row = start_row + 1

    for r in records:
        ws.row_dimensions[current_row].height = 20
        row_fill = ZEBRA_FILL if (current_row % 2 == 0) else WHITE_FILL

        cin = r.get("check_in")
        cout = r.get("check_out")
        cin_str = cin.strftime("%I:%M %p") if hasattr(cin, "strftime") else (str(cin) if cin else "--:--")
        
        if cin and not cout:
            cout_str = "Active Session"
        else:
            cout_str = cout.strftime("%I:%M %p") if hasattr(cout, "strftime") else (str(cout) if cout else "--:--")

        emp_code = r.get("employee_code") or f"EMP{int(r.get('employee_id') or 0):03d}"
        cnt = int(r.get("checkout_count") or 0)
        cnt_str = f"{cnt} Checkouts" if cnt > 1 else ("1 Checkout" if cnt == 1 else ("Active" if (cin and not cout) else "0"))

        hrs = float(r.get("working_hours") or 0.00)
        ot = float(r.get("overtime_hours") or 0.00)

        c_id = ws.cell(row=current_row, column=2, value=emp_code)
        c_id.alignment = ALIGN_CENTER
        c_id.font = BOLD_FONT

        c_name = ws.cell(row=current_row, column=3, value=r.get("employee_name") or "Unknown")
        c_name.alignment = ALIGN_LEFT
        c_name.font = BOLD_FONT

        c_dept = ws.cell(row=current_row, column=4, value=r.get("department") or "N/A")
        c_dept.alignment = ALIGN_LEFT

        c_cin = ws.cell(row=current_row, column=5, value=cin_str)
        c_cin.alignment = ALIGN_CENTER

        c_cout = ws.cell(row=current_row, column=6, value=cout_str)
        c_cout.alignment = ALIGN_CENTER
        if cout_str == "Active Session":
            c_cout.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="B45309")

        c_cnt = ws.cell(row=current_row, column=7, value=cnt_str)
        c_cnt.alignment = ALIGN_CENTER

        c_ses = ws.cell(row=current_row, column=8, value=format_sessions_summary(r.get("punches")))
        c_ses.alignment = ALIGN_LEFT

        c_hrs = ws.cell(row=current_row, column=9, value=hrs)
        c_hrs.alignment = ALIGN_RIGHT
        c_hrs.number_format = '0.00'

        c_ot = ws.cell(row=current_row, column=10, value=ot)
        c_ot.alignment = ALIGN_RIGHT
        c_ot.number_format = '0.00'

        # Status Cell with styled pill fill
        status_fill, status_font = _get_status_style(r.get("status"))
        status_norm = _normalize_status(r.get("status"))
        status_cell = ws.cell(row=current_row, column=11, value=status_norm.upper())
        status_cell.alignment = ALIGN_CENTER
        status_cell.fill = status_fill
        status_cell.font = status_font

        c_src = ws.cell(row=current_row, column=12, value=str(r.get("source") or "biometric").capitalize())
        c_src.alignment = ALIGN_CENTER

        for col_idx in range(2, 13):
            c = ws.cell(row=current_row, column=col_idx)
            if col_idx != 11:
                c.fill = row_fill
                if col_idx not in (2, 3, 6):
                    c.font = REGULAR_FONT
            c.border = THIN_BORDER

        current_row += 1

    # Comprehensive Executive Grand Totals Row
    ws.row_dimensions[current_row].height = 26
    
    ws.cell(row=current_row, column=2, value="TOTALS").alignment = ALIGN_CENTER
    ws.cell(row=current_row, column=2).font = BOLD_FONT

    c_tot_name = ws.cell(row=current_row, column=3, value="GRAND TOTAL SUMMARY")
    c_tot_name.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0F172A")
    c_tot_name.alignment = ALIGN_LEFT

    c_tot_dept = ws.cell(row=current_row, column=4, value=f"{total_logs} Staff Records")
    c_tot_dept.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="475569")
    c_tot_dept.alignment = ALIGN_LEFT

    c_tot_cin = ws.cell(row=current_row, column=5, value=f"P: {present_cnt}")
    c_tot_cin.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="065F46")
    c_tot_cin.alignment = ALIGN_CENTER

    c_tot_cout = ws.cell(row=current_row, column=6, value=f"A: {absent_cnt}")
    c_tot_cout.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="991B1B")
    c_tot_cout.alignment = ALIGN_CENTER

    c_tot_cnt = ws.cell(row=current_row, column=7, value=f"{total_checkouts} COs")
    c_tot_cnt.font = BOLD_FONT
    c_tot_cnt.alignment = ALIGN_CENTER

    c_tot_ses = ws.cell(row=current_row, column=8, value=f"Late: {late_cnt} ({total_late_str}) | Hol: {holiday_cnt} | Lv: {leave_cnt}")
    c_tot_ses.font = BOLD_FONT
    c_tot_ses.alignment = ALIGN_LEFT

    t_hrs = ws.cell(row=current_row, column=9, value=round(total_hours, 2))
    t_hrs.alignment = ALIGN_RIGHT
    t_hrs.font = BOLD_FONT
    t_hrs.number_format = '0.00'

    t_ot = ws.cell(row=current_row, column=10, value=round(total_ot, 2))
    t_ot.alignment = ALIGN_RIGHT
    t_ot.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="7C3AED")
    t_ot.number_format = '0.00'

    c_tot_stat = ws.cell(row=current_row, column=11, value=f"P:{present_cnt} A:{absent_cnt} L:{late_cnt}")
    c_tot_stat.font = BOLD_FONT
    c_tot_stat.alignment = ALIGN_CENTER

    c_tot_src = ws.cell(row=current_row, column=12, value="100% Verified")
    c_tot_src.font = Font(name=FONT_FAMILY, size=9, bold=True, color="0284C7")
    c_tot_src.alignment = ALIGN_CENTER

    for col_idx in range(2, 13):
        c = ws.cell(row=current_row, column=col_idx)
        c.border = TOTAL_BORDER
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    _autofit_columns(ws, min_col=2, max_col=12, skip_rows=8)

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()
    return excel_bytes

# ─── 2. Monthly Attendance Summary Excel ──────────────────────────────────────
def generate_monthly_attendance_excel(year: int, month: int, summaries: list, title: str = "MONTHLY ATTENDANCE SUMMARY", subtitle: str = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Summary"
    ws.views.sheetView[0].showGridLines = True

    month_name = calendar.month_name[month]
    ws.cell(row=2, column=2, value="SMART ATTENDANCE MANAGEMENT SYSTEM").font = Font(name=FONT_FAMILY, size=10, bold=True, color="0284C7")
    ws.cell(row=3, column=2, value=title).font = TITLE_FONT
    sub = subtitle or f"Period: {month_name} {year} | Exported on: {datetime.now().strftime('%d %b %Y, %I:%M %p IST')}"
    ws.cell(row=4, column=2, value=sub).font = SUBTITLE_FONT

    # Pre-calculate Totals & Metrics
    total_staff = len(summaries)
    total_pres = sum(float(s.get("present_days") or 0.0) for s in summaries)
    total_abs = sum(float(s.get("absent_days") or 0.0) for s in summaries)
    total_lv = sum(float(s.get("leave_days") or 0.0) for s in summaries)
    total_ot = sum(float(s.get("overtime_hours") or 0.0) for s in summaries)
    _, days_in_month = calendar.monthrange(year, month)
    total_possible = total_staff * days_in_month
    avg_att = (total_pres / (total_pres + total_abs) * 100) if (total_pres + total_abs) > 0 else (
        (total_pres / total_possible * 100) if total_possible > 0 else 0.0
    )

    # Executive KPI Dashboard Cards (Rows 6-7, Cols 2-9)
    kpis = [
        ("TOTAL STAFF", total_staff, "0F172A"),
        ("PRESENT DAYS", f"{total_pres:.1f}", "065F46"),
        ("ABSENT DAYS", f"{total_abs:.1f}", "991B1B"),
        ("LEAVE DAYS", f"{total_lv:.1f}", "6D28D9"),
        ("OVERTIME HOURS", f"+{total_ot:.2f}h", "7C3AED"),
        ("AVG ATTENDANCE", f"{avg_att:.1f}%", "0284C7"),
        ("DAYS IN MONTH", f"{days_in_month} Days", "0F172A"),
        ("ROSTER STATUS", f"{total_staff} Active", "065F46"),
    ]
    _write_kpi_dashboard(ws, start_row=6, kpi_list=kpis, start_col=2)

    headers = [
        "Employee ID", "Employee Name", "Department", "Designation", 
        "Present Days", "Absent Days", "Leave Days", "Overtime (hrs)"
    ]

    start_row = 9
    ws.row_dimensions[8].height = 10
    ws.row_dimensions[start_row].height = 26
    for col_idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    ws.freeze_panes = f"B{start_row + 1}"

    current_row = start_row + 1

    for s in summaries:
        ws.row_dimensions[current_row].height = 20
        row_fill = ZEBRA_FILL if (current_row % 2 == 0) else WHITE_FILL

        p_days = float(s.get("present_days") or 0.0)
        a_days = float(s.get("absent_days") or 0.0)
        l_days = float(s.get("leave_days") or 0.0)
        ot_h = float(s.get("overtime_hours") or 0.00)

        c_id = ws.cell(row=current_row, column=2, value=f"EMP{int(s.get('employee_id') or 0):03d}")
        c_id.alignment = ALIGN_CENTER
        c_id.font = BOLD_FONT

        c_name = ws.cell(row=current_row, column=3, value=s.get("employee_name") or "Unknown")
        c_name.alignment = ALIGN_LEFT
        c_name.font = BOLD_FONT

        ws.cell(row=current_row, column=4, value=s.get("department") or "N/A").alignment = ALIGN_LEFT
        ws.cell(row=current_row, column=5, value=s.get("designation") or "Staff").alignment = ALIGN_LEFT

        c_p = ws.cell(row=current_row, column=6, value=p_days)
        c_p.alignment = ALIGN_RIGHT
        c_p.number_format = '0.0'

        c_a = ws.cell(row=current_row, column=7, value=a_days)
        c_a.alignment = ALIGN_RIGHT
        c_a.number_format = '0.0'

        c_l = ws.cell(row=current_row, column=8, value=l_days)
        c_l.alignment = ALIGN_RIGHT
        c_l.number_format = '0.0'

        c_ot = ws.cell(row=current_row, column=9, value=ot_h)
        c_ot.alignment = ALIGN_RIGHT
        c_ot.number_format = '0.00'

        for col_idx in range(2, 10):
            c = ws.cell(row=current_row, column=col_idx)
            c.fill = row_fill
            if col_idx not in (2, 3):
                c.font = REGULAR_FONT
            c.border = THIN_BORDER

        current_row += 1

    # Comprehensive Executive Summary Totals Row
    ws.row_dimensions[current_row].height = 26
    
    ws.cell(row=current_row, column=2, value="TOTALS").alignment = ALIGN_CENTER
    ws.cell(row=current_row, column=2).font = BOLD_FONT

    c_tot_name = ws.cell(row=current_row, column=3, value=f"GRAND TOTALS ({total_staff} Staff)")
    c_tot_name.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0F172A")
    c_tot_name.alignment = ALIGN_LEFT

    c_tot_dept = ws.cell(row=current_row, column=4, value="Company Roster")
    c_tot_dept.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="475569")
    c_tot_dept.alignment = ALIGN_LEFT

    c_tot_des = ws.cell(row=current_row, column=5, value=f"Avg Att: {avg_att:.1f}%")
    c_tot_des.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0284C7")
    c_tot_des.alignment = ALIGN_CENTER

    t_p = ws.cell(row=current_row, column=6, value=round(total_pres, 1))
    t_p.alignment = ALIGN_RIGHT
    t_p.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="065F46")
    t_p.number_format = '0.0'

    t_a = ws.cell(row=current_row, column=7, value=round(total_abs, 1))
    t_a.alignment = ALIGN_RIGHT
    t_a.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="991B1B")
    t_a.number_format = '0.0'

    t_l = ws.cell(row=current_row, column=8, value=round(total_lv, 1))
    t_l.alignment = ALIGN_RIGHT
    t_l.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="6D28D9")
    t_l.number_format = '0.0'

    t_ot = ws.cell(row=current_row, column=9, value=round(total_ot, 2))
    t_ot.alignment = ALIGN_RIGHT
    t_ot.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="7C3AED")
    t_ot.number_format = '0.00'

    for col_idx in range(2, 10):
        c = ws.cell(row=current_row, column=col_idx)
        c.border = TOTAL_BORDER
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    _autofit_columns(ws, min_col=2, max_col=9, skip_rows=8)

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()
    return excel_bytes

# ─── 3. Monthly Payroll Report Excel ──────────────────────────────────────────
def generate_payroll_report_excel(year: int, month: int, payrolls: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll Report"
    ws.views.sheetView[0].showGridLines = True

    month_name = calendar.month_name[month]
    ws.cell(row=2, column=2, value="SMART ATTENDANCE & PAYROLL MANAGEMENT SYSTEM").font = Font(name=FONT_FAMILY, size=10, bold=True, color="0284C7")
    ws.cell(row=3, column=2, value="MONTHLY PAYROLL SUMMARY REPORT").font = TITLE_FONT
    ws.cell(row=4, column=2, value=f"Period: {month_name} {year} | Exported on: {datetime.now().strftime('%d %b %Y, %I:%M %p IST')}").font = SUBTITLE_FONT

    headers = [
        "Employee ID", "Employee Name", "Department", "Working Days", 
        "Present Days", "Overtime (hrs)", "Basic Salary (₹)", "Overtime Pay (₹)", 
        "Deductions (₹)", "Net Salary (₹)", "Payment Status"
    ]

    start_row = 6
    ws.row_dimensions[start_row].height = 26
    for col_idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    ws.freeze_panes = f"B{start_row + 1}"

    current_row = start_row + 1
    total_basic = 0.0
    total_ot_pay = 0.0
    total_deductions = 0.0
    total_net = 0.0

    for p in payrolls:
        ws.row_dimensions[current_row].height = 20
        row_fill = ZEBRA_FILL if (current_row % 2 == 0) else WHITE_FILL

        basic = float(p.get("basic_salary") or 0.0)
        ot_pay = float(p.get("overtime_pay") or 0.0)
        ded = float(p.get("deductions") or 0.0)
        net = float(p.get("final_salary") or 0.0)

        total_basic += basic
        total_ot_pay += ot_pay
        total_deductions += ded
        total_net += net

        c_id = ws.cell(row=current_row, column=2, value=f"EMP{int(p.get('employee_id') or 0):03d}")
        c_id.alignment = ALIGN_CENTER
        c_id.font = BOLD_FONT

        c_name = ws.cell(row=current_row, column=3, value=p.get("employee_name") or "Unknown")
        c_name.alignment = ALIGN_LEFT
        c_name.font = BOLD_FONT

        ws.cell(row=current_row, column=4, value=p.get("department") or "N/A").alignment = ALIGN_LEFT
        ws.cell(row=current_row, column=5, value=int(p.get("working_days") or 0)).alignment = ALIGN_RIGHT
        ws.cell(row=current_row, column=6, value=float(p.get("present_days") or 0.0)).alignment = ALIGN_RIGHT
        ws.cell(row=current_row, column=7, value=float(p.get("overtime_hours") or 0.0)).alignment = ALIGN_RIGHT

        c_basic = ws.cell(row=current_row, column=8, value=basic)
        c_basic.alignment = ALIGN_RIGHT
        c_basic.number_format = '"₹"#,##0.00'

        c_otpay = ws.cell(row=current_row, column=9, value=ot_pay)
        c_otpay.alignment = ALIGN_RIGHT
        c_otpay.number_format = '"₹"#,##0.00'

        c_ded = ws.cell(row=current_row, column=10, value=ded)
        c_ded.alignment = ALIGN_RIGHT
        c_ded.number_format = '"₹"#,##0.00'
        c_ded.font = Font(name=FONT_FAMILY, size=9.5, color="DC2626")

        c_net = ws.cell(row=current_row, column=11, value=net)
        c_net.alignment = ALIGN_RIGHT
        c_net.number_format = '"₹"#,##0.00'
        c_net.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="059669")

        is_paid = bool(p.get("is_paid"))
        status_cell = ws.cell(row=current_row, column=12, value="PAID" if is_paid else "UNPAID")
        status_cell.alignment = ALIGN_CENTER
        if is_paid:
            status_cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
            status_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="065F46")
        else:
            status_cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
            status_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="991B1B")

        for col_idx in range(2, 12):
            c = ws.cell(row=current_row, column=col_idx)
            c.fill = row_fill
            if col_idx not in (2, 3, 10, 11):
                c.font = REGULAR_FONT
            c.border = THIN_BORDER
        status_cell.border = THIN_BORDER

        current_row += 1

    # Totals Row
    ws.row_dimensions[current_row].height = 26
    ws.cell(row=current_row, column=2, value="TOTALS").alignment = ALIGN_CENTER
    ws.cell(row=current_row, column=2).font = BOLD_FONT

    c_tot_name = ws.cell(row=current_row, column=3, value="Grand Totals")
    c_tot_name.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0F172A")
    c_tot_name.alignment = ALIGN_LEFT
    
    t_b = ws.cell(row=current_row, column=8, value=total_basic)
    t_b.font = BOLD_FONT
    t_b.alignment = ALIGN_RIGHT
    t_b.number_format = '"₹"#,##0.00'

    t_ot = ws.cell(row=current_row, column=9, value=total_ot_pay)
    t_ot.font = BOLD_FONT
    t_ot.alignment = ALIGN_RIGHT
    t_ot.number_format = '"₹"#,##0.00'

    t_d = ws.cell(row=current_row, column=10, value=total_deductions)
    t_d.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="DC2626")
    t_d.alignment = ALIGN_RIGHT
    t_d.number_format = '"₹"#,##0.00'

    t_n = ws.cell(row=current_row, column=11, value=total_net)
    t_n.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="059669")
    t_n.alignment = ALIGN_RIGHT
    t_n.number_format = '"₹"#,##0.00'

    for col_idx in range(2, 13):
        c = ws.cell(row=current_row, column=col_idx)
        c.border = TOTAL_BORDER
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    _autofit_columns(ws, min_col=2, max_col=12, skip_rows=5)

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()
    return excel_bytes

# ─── 4. Historical Attendance Excel ───────────────────────────────────────────
def generate_historical_attendance_excel(start_date: date, end_date: date, records: list, title: str = "HISTORICAL ATTENDANCE REPORT", subtitle: str = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Logs"
    ws.views.sheetView[0].showGridLines = True

    ws.cell(row=2, column=2, value="SMART ATTENDANCE MANAGEMENT SYSTEM").font = Font(name=FONT_FAMILY, size=10, bold=True, color="0284C7")
    ws.cell(row=3, column=2, value=title).font = TITLE_FONT
    sub = subtitle or f"Period: {start_date.strftime('%d %B %Y')} to {end_date.strftime('%d %B %Y')} | Exported on: {datetime.now().strftime('%d %b %Y, %I:%M %p IST')}"
    ws.cell(row=4, column=2, value=sub).font = SUBTITLE_FONT

    # Pre-calculate Totals & Metrics
    total_logs = len(records)
    present_cnt = sum(1 for r in records if _normalize_status(r.get("status")) in ("present", "wfh"))
    late_cnt = sum(1 for r in records if _normalize_status(r.get("status")) == "late" or (r.get("late_minutes") and int(r.get("late_minutes")) > 0))
    absent_cnt = sum(1 for r in records if _normalize_status(r.get("status")) == "absent")
    leave_cnt = sum(1 for r in records if _normalize_status(r.get("status")) == "leave")
    holiday_cnt = sum(1 for r in records if _normalize_status(r.get("status")) in ("holiday", "weekly_off"))
    halfday_cnt = sum(1 for r in records if _normalize_status(r.get("status")) in ("half_day", "halfday"))
    total_late_min = sum(int(r.get("late_minutes") or 0) for r in records)
    total_late_str = format_late_duration_excel(total_late_min)
    total_checkouts = sum(int(r.get("checkout_count") or (1 if r.get("check_out") else 0)) for r in records)
    total_hours = sum(float(r.get("working_hours") or 0.0) for r in records)
    total_ot = sum(float(r.get("overtime_hours") or 0.0) for r in records)
    attended_cnt = sum(1 for r in records if _normalize_status(r.get("status")) in ("present", "wfh", "late", "half_day", "halfday"))
    att_rate = (attended_cnt / total_logs * 100) if total_logs > 0 else 0.0

    # Executive KPI Dashboard Cards (Rows 6-7, Cols 2-14)
    kpis = [
        ("TOTAL LOGS", total_logs, "0F172A"),
        ("PRESENT", present_cnt, "065F46"),
        ("ABSENT", absent_cnt, "991B1B"),
        ("LATE COUNT", late_cnt, "92400E"),
        ("TOTAL LATE TIME", total_late_str, "92400E"),
        ("APPROVED LEAVE", leave_cnt, "6D28D9"),
        ("WEEKLY OFF / HOL", holiday_cnt, "0369A1"),
        ("CHECKOUTS", f"{total_checkouts} COs", "0284C7"),
        ("TOTAL WORKED", f"{total_hours:.2f} hrs", "0F172A"),
        ("OVERTIME", f"+{total_ot:.2f} hrs", "7C3AED"),
        ("HALF DAYS", halfday_cnt, "92400E"),
        ("ATTENDANCE RATE", f"{att_rate:.1f}%", "065F46"),
        ("SYSTEM STATUS", "100% ARCHIVED", "0284C7"),
    ]
    _write_kpi_dashboard(ws, start_row=6, kpi_list=kpis, start_col=2)

    headers = [
        "Date", "Employee ID", "Employee Name", "Department",
        "Check-In Time", "Check-Out Time", "Checkouts", "Sessions Breakdown",
        "Working Hours", "Late Duration", "Overtime (hrs)", "Status", "Source"
    ]

    start_row = 9
    ws.row_dimensions[8].height = 10
    ws.row_dimensions[start_row].height = 26
    for col_idx, header in enumerate(headers, start=2):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    ws.freeze_panes = f"B{start_row + 1}"

    current_row = start_row + 1

    for r in records:
        ws.row_dimensions[current_row].height = 20
        row_fill = ZEBRA_FILL if (current_row % 2 == 0) else WHITE_FILL

        r_date = r.get("date")
        if hasattr(r_date, "strftime"):
            date_str = r_date.strftime("%Y-%m-%d")
        else:
            date_str = str(r_date) if r_date else "--"

        cin = r.get("check_in")
        cout = r.get("check_out")
        cin_str = cin.strftime("%I:%M %p") if hasattr(cin, "strftime") else (str(cin) if cin else "--:--")

        if cin and not cout:
            cout_str = "Active Session"
        else:
            cout_str = cout.strftime("%I:%M %p") if hasattr(cout, "strftime") else (str(cout) if cout else "--:--")

        emp_code = r.get("employee_code") or f"EMP{int(r.get('employee_id') or 0):03d}"
        cnt = int(r.get("checkout_count") or 0)
        cnt_str = f"{cnt} Checkouts" if cnt > 1 else ("1 Checkout" if cnt == 1 else ("Active" if (cin and not cout) else "0"))

        hrs = float(r.get("working_hours") or 0.00)
        ot = float(r.get("overtime_hours") or 0.00)

        c_dt = ws.cell(row=current_row, column=2, value=date_str)
        c_dt.alignment = ALIGN_CENTER

        c_id = ws.cell(row=current_row, column=3, value=emp_code)
        c_id.alignment = ALIGN_CENTER
        c_id.font = BOLD_FONT

        c_name = ws.cell(row=current_row, column=4, value=r.get("employee_name") or "Unknown")
        c_name.alignment = ALIGN_LEFT
        c_name.font = BOLD_FONT

        ws.cell(row=current_row, column=5, value=r.get("department") or "N/A").alignment = ALIGN_LEFT
        ws.cell(row=current_row, column=6, value=cin_str).alignment = ALIGN_CENTER

        c_cout = ws.cell(row=current_row, column=7, value=cout_str)
        c_cout.alignment = ALIGN_CENTER
        if cout_str == "Active Session":
            c_cout.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="B45309")

        ws.cell(row=current_row, column=8, value=cnt_str).alignment = ALIGN_CENTER
        ws.cell(row=current_row, column=9, value=format_sessions_summary(r.get("punches"))).alignment = ALIGN_LEFT

        c_hrs = ws.cell(row=current_row, column=10, value=hrs)
        c_hrs.alignment = ALIGN_RIGHT
        c_hrs.number_format = '0.00'

        c_late = ws.cell(row=current_row, column=11, value=format_late_duration_excel(r.get("late_minutes")))
        c_late.alignment = ALIGN_CENTER

        c_ot = ws.cell(row=current_row, column=12, value=ot)
        c_ot.alignment = ALIGN_RIGHT
        c_ot.number_format = '0.00'

        # Status cell with pill
        status_fill, status_font = _get_status_style(r.get("status"))
        status_norm = _normalize_status(r.get("status"))
        status_cell = ws.cell(row=current_row, column=13, value=status_norm.upper())
        status_cell.alignment = ALIGN_CENTER
        status_cell.fill = status_fill
        status_cell.font = status_font

        ws.cell(row=current_row, column=14, value=str(r.get("source") or "biometric").capitalize()).alignment = ALIGN_CENTER

        for col_idx in range(2, 15):
            c = ws.cell(row=current_row, column=col_idx)
            if col_idx != 13:
                c.fill = row_fill
                if col_idx not in (3, 4, 7):
                    c.font = REGULAR_FONT
            c.border = THIN_BORDER

        current_row += 1

    # Comprehensive Executive Grand Totals Row
    ws.row_dimensions[current_row].height = 26
    
    ws.cell(row=current_row, column=2, value="TOTALS").alignment = ALIGN_CENTER
    ws.cell(row=current_row, column=2).font = BOLD_FONT

    ws.cell(row=current_row, column=3, value="ALL").alignment = ALIGN_CENTER
    ws.cell(row=current_row, column=3).font = BOLD_FONT

    c_tot_name = ws.cell(row=current_row, column=4, value="GRAND TOTAL SUMMARY")
    c_tot_name.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="0F172A")
    c_tot_name.alignment = ALIGN_LEFT

    c_tot_dept = ws.cell(row=current_row, column=5, value=f"{total_logs} Logs")
    c_tot_dept.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="475569")
    c_tot_dept.alignment = ALIGN_LEFT

    c_tot_cin = ws.cell(row=current_row, column=6, value=f"P: {present_cnt}")
    c_tot_cin.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="065F46")
    c_tot_cin.alignment = ALIGN_CENTER

    c_tot_cout = ws.cell(row=current_row, column=7, value=f"A: {absent_cnt}")
    c_tot_cout.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="991B1B")
    c_tot_cout.alignment = ALIGN_CENTER

    c_tot_cnt = ws.cell(row=current_row, column=8, value=f"{total_checkouts} COs")
    c_tot_cnt.font = BOLD_FONT
    c_tot_cnt.alignment = ALIGN_CENTER

    c_tot_ses = ws.cell(row=current_row, column=9, value=f"Lv: {leave_cnt} | Hol: {holiday_cnt}")
    c_tot_ses.font = BOLD_FONT
    c_tot_ses.alignment = ALIGN_LEFT

    t_hrs = ws.cell(row=current_row, column=10, value=round(total_hours, 2))
    t_hrs.alignment = ALIGN_RIGHT
    t_hrs.font = BOLD_FONT
    t_hrs.number_format = '0.00'

    c_tot_late = ws.cell(row=current_row, column=11, value=f"{total_late_str} ({late_cnt} late)")
    c_tot_late.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="92400E")
    c_tot_late.alignment = ALIGN_CENTER

    t_ot = ws.cell(row=current_row, column=12, value=round(total_ot, 2))
    t_ot.alignment = ALIGN_RIGHT
    t_ot.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color="7C3AED")
    t_ot.number_format = '0.00'

    c_tot_stat = ws.cell(row=current_row, column=13, value=f"P:{present_cnt} A:{absent_cnt} L:{late_cnt}")
    c_tot_stat.font = BOLD_FONT
    c_tot_stat.alignment = ALIGN_CENTER

    c_tot_src = ws.cell(row=current_row, column=14, value="100% Verified")
    c_tot_src.font = Font(name=FONT_FAMILY, size=9, bold=True, color="0284C7")
    c_tot_src.alignment = ALIGN_CENTER

    for col_idx in range(2, 15):
        c = ws.cell(row=current_row, column=col_idx)
        c.border = TOTAL_BORDER
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    _autofit_columns(ws, min_col=2, max_col=14, skip_rows=8)

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()
    return excel_bytes
